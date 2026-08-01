#!/usr/bin/env python
"""Synchronize CIL_Baselines.xlsx from complete, canonically named logs.

Only a single log with a complete CNN curve may update a blank workbook result
pair. Existing values are retained and reported as matches or mismatches.
"""

from __future__ import annotations

import ast
import csv
import os
import re
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from utils.logging_utils import canonical_log_dir


WORKBOOK = ROOT / "CIL_Baselines.xlsx"
CSV_REPORT = ROOT / "baseline_log_mapping.csv"
LOG_ROOT = ROOT / "logs"

METHOD_MODELS = {
    "Full Fine-Tuning": "finetune",
    "SimpleCIL": "simplecil",
    "APER": "adam_adapter",
    "L2P": "l2p",
    "DualPrompt": "dualprompt",
    "CODA-Prompt": "coda_prompt",
    "LAE": "lae",
    "InfLoRA": "inflora",
    "EASE": "ease",
    "BiLoRA": "bilora",
    "SD-LoRA": "sdlora",
    "CL-LoRA": "cllora",
    "iCaRL": "icarl",
    "FOSTER": "foster",
}

DATASETS = (
    ("CIFAR100", ("C", "D"), ("cifar224",), 100, 10),
    ("CUB200", ("E", "F"), ("cub",), 200, 10),
    ("ImageNet-R", ("G", "H"), ("imagenetr", "imagenet_r"), 200, 5),
    ("OmniBenchmark", ("I", "J"), ("omnibenchmark",), 300, 10),
    ("VTAB", ("K", "L"), ("vtab",), 50, 5),
)

CURVE_RE = re.compile(r"CNN top1 curve: (\[[^\r\n]+\])")
AVERAGE_RE = re.compile(r"Average Accuracy \(CNN\): ([0-9.]+)")


def compact(value):
    return "" if value is None else str(value)


def parse_log(path: Path, expected_tasks: int):
    text = path.read_text(errors="ignore")
    curves = CURVE_RE.findall(text)
    averages = AVERAGE_RE.findall(text)
    if not curves:
        return {"status": "missing_curve", "al": "", "average": ""}
    try:
        curve = ast.literal_eval(curves[-1])
    except (SyntaxError, ValueError):
        return {"status": "unparseable_curve", "al": "", "average": ""}
    if len(curve) != expected_tasks:
        return {"status": "incomplete_{}of{}".format(len(curve), expected_tasks), "al": "", "average": ""}
    if not averages:
        return {"status": "missing_average", "al": "", "average": ""}
    return {"status": "complete", "al": float(curve[-1]), "average": float(averages[-1])}


def candidates(method: str, prefixes, increment: int, seed: int):
    directory = LOG_ROOT / canonical_log_dir(METHOD_MODELS[method])
    suffix = "_0_{}_{}.log".format(increment, seed).lower()
    if not directory.is_dir():
        return []
    return sorted(
        path
        for path in directory.glob("*.log")
        if path.name.lower().endswith(suffix)
        and any(path.name.lower().startswith(prefix) for prefix in prefixes)
    )


def compare(existing_al, existing_average, details):
    if len(details) != 1 or details[0]["status"] != "complete":
        return "not_verifiable"
    if existing_al == "" and existing_average == "":
        return "fill"
    if existing_al == "" or existing_average == "":
        return "partial_workbook_pair"
    detail = details[0]
    try:
        same_al = abs(float(existing_al) - detail["al"]) < 0.005
        same_average = abs(float(existing_average) - detail["average"]) < 0.005
    except (TypeError, ValueError):
        return "non_numeric_workbook_value"
    return "match" if same_al and same_average else "mismatch"


def main():
    if not WORKBOOK.is_file():
        raise SystemExit("Missing workbook: {}".format(WORKBOOK))
    workbook = load_workbook(WORKBOOK)
    rows = []
    updates = {}

    for sheet in workbook.worksheets:
        if not sheet.title.startswith("seed"):
            continue
        try:
            seed = int(sheet.title.removeprefix("seed"))
        except ValueError:
            continue
        for row_number in range(3, sheet.max_row + 1):
            method = sheet.cell(row_number, 1).value
            if method not in METHOD_MODELS:
                continue
            for dataset, columns, prefixes, class_count, task_count in DATASETS:
                increment = class_count // task_count
                paths = candidates(method, prefixes, increment, seed)
                details = [parse_log(path, task_count) for path in paths]
                existing_al = compact(sheet["{}{}".format(columns[0], row_number)].value)
                existing_average = compact(sheet["{}{}".format(columns[1], row_number)].value)
                action = compare(existing_al, existing_average, details)
                if action == "fill":
                    detail = details[0]
                    updates[(sheet.title, "{}{}".format(columns[0], row_number))] = detail["al"]
                    updates[(sheet.title, "{}{}".format(columns[1], row_number))] = detail["average"]
                status = "missing_log" if not paths else ";".join(detail["status"] for detail in details)
                rows.append({
                    "sheet": sheet.title,
                    "seed": seed,
                    "method": method,
                    "dataset": dataset,
                    "cells": "{}{}:{}{}".format(columns[0], row_number, columns[1], row_number),
                    "paths": ";".join(str(path) for path in paths),
                    "log_al": ";".join(str(detail["al"]) for detail in details),
                    "log_average": ";".join(str(detail["average"]) for detail in details),
                    "status": status,
                    "workbook_action": action,
                })

    if updates:
        temporary = WORKBOOK.with_name(WORKBOOK.stem + ".validation-copy.xlsx")
        for (sheet_name, cell), value in updates.items():
            if workbook[sheet_name][cell].value not in (None, ""):
                raise RuntimeError("Refusing to overwrite {}!{}".format(sheet_name, cell))
            workbook[sheet_name][cell] = value
        workbook.save(temporary)
        verified = load_workbook(temporary, read_only=True, data_only=True)
        for (sheet_name, cell), expected in updates.items():
            actual = verified[sheet_name][cell].value
            if actual is None or abs(float(actual) - expected) > 1e-9:
                raise RuntimeError("Validation failed for {}!{}".format(sheet_name, cell))
        os.replace(temporary, WORKBOOK)

    with CSV_REPORT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=tuple(rows[0]) if rows else ())
        if rows:
            writer.writeheader()
            writer.writerows(rows)

    print("updated_cells={}".format(len(updates)))
    print("log_status={}".format(dict(sorted(Counter(row["status"] for row in rows).items()))))
    print("workbook_action={}".format(dict(sorted(Counter(row["workbook_action"] for row in rows).items()))))


if __name__ == "__main__":
    main()
