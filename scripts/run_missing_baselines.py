#!/usr/bin/env python
"""Run blank CIL_Baselines.xlsx entries on one explicitly selected GPU."""

from __future__ import annotations

import ast
import json
import os
import re
import sys
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from trainer import train
from utils.logging_utils import canonical_log_dir


WORKBOOK = ROOT / "CIL_Baselines.xlsx"
DATA_ROOT = Path("/home/team/zhaohongwei/Dataset")
RUN_OUTPUT = ROOT / "run_output"

METHOD_CONFIGS = {
    "Full Fine-Tuning": ("finetune", "finetune"),
    "SimpleCIL": ("simplecil", "simplecil"),
    "APER": ("aper", "aper_adapter"),
    "L2P": ("l2p", "l2p"),
    "DualPrompt": ("dualprompt", "dualprompt"),
    "CODA-Prompt": ("coda_prompt", "coda_prompt"),
    "LAE": ("lae", "lae"),
    "InfLoRA": ("inflora", "inflora"),
    "EASE": ("ease", "ease"),
    "BiLoRA": ("bilora", "bilora"),
    "SD-LoRA": ("sdlora", "sdlora"),
    "CL-LoRA": ("cllora", "cllora"),
    "iCaRL": ("icarl", "icarl"),
    "FOSTER": ("foster", "foster"),
}

DATASETS = (
    ("CIFAR100", ("C", "D"), ("cifar",), "cifar224", 100, 10),
    ("CUB200", ("E", "F"), ("cub",), "cub", 200, 20),
    ("ImageNet-R", ("G", "H"), ("inr",), "imagenetr", 200, 40),
    # Both abbreviations are present in the upstream configuration tree.
    ("OmniBenchmark", ("I", "J"), ("omn", "omni"), "omnibenchmark", 300, 30),
    ("VTAB", ("K", "L"), ("vtab",), "vtab", 50, 10),
)

CURVE_RE = re.compile(r"CNN top1 curve: (\[[^\r\n]+\])")
AVERAGE_RE = re.compile(r"Average Accuracy \(CNN\): ([0-9.]+)")


def is_complete(path: Path, expected_tasks: int):
    if not path.is_file():
        return False
    text = path.read_text(errors="ignore")
    curves = CURVE_RE.findall(text)
    if not curves or not AVERAGE_RE.findall(text):
        return False
    try:
        return len(ast.literal_eval(curves[-1])) == expected_tasks
    except (SyntaxError, ValueError):
        return False


def build_tasks():
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    tasks = []
    for sheet in workbook.worksheets:
        if not sheet.title.startswith("seed"):
            continue
        try:
            seed = int(sheet.title.removeprefix("seed"))
        except ValueError:
            continue
        for row in range(3, sheet.max_row + 1):
            method = sheet.cell(row, 1).value
            if method not in METHOD_CONFIGS:
                continue
            for dataset, columns, config_datasets, data_key, class_count, increment in DATASETS:
                values = (sheet["{}{}".format(columns[0], row)].value, sheet["{}{}".format(columns[1], row)].value)
                if values != (None, None):
                    continue
                config_dir, config_prefix = METHOD_CONFIGS[method]
                candidates = [
                    ROOT / "configs" / config_dir / "{}_{}_B0_Inc{}.json".format(config_prefix, config_dataset, increment)
                    for config_dataset in config_datasets
                ]
                config = next((candidate for candidate in candidates if candidate.is_file()), None)
                if config is None:
                    print(
                        "MISSING_CONFIG method={} dataset={} paths={}".format(
                            method, dataset, ",".join(str(candidate) for candidate in candidates)
                        ),
                        flush=True,
                    )
                    continue
                tasks.append({
                    "sheet": sheet.title,
                    "seed": seed,
                    "method": method,
                    "dataset": dataset,
                    "config": config,
                    "data_key": data_key,
                    "class_count": class_count,
                    "increment": increment,
                })
    return tasks


def log_path(args, seed):
    backbone = args["backbone_type"].removeprefix("pretrained_")
    init_cls = 0 if args["init_cls"] == args["increment"] else args["init_cls"]
    filename = "{}_{}_{}_{}_{}.log".format(args["dataset"], backbone, init_cls, args["increment"], seed)
    return ROOT / "logs" / canonical_log_dir(args["model_name"]) / filename


def run_task(task, gpu_id):
    with task["config"].open() as handle:
        args = json.load(handle)
    expected_tasks = task["class_count"] // task["increment"]
    output = log_path(args, task["seed"])
    if is_complete(output, expected_tasks):
        print("SKIP complete {}".format(output), flush=True)
        return True

    lock_root = RUN_OUTPUT / "locks"
    lock_root.mkdir(parents=True, exist_ok=True)
    lock = lock_root / "{}_{}_{}_{}.lock".format(task["sheet"], task["method"].replace(" ", "_"), task["dataset"], task["seed"])
    try:
        lock.mkdir()
    except FileExistsError:
        pid_path = lock / "pid"
        try:
            pid = int(pid_path.read_text().strip())
            os.kill(pid, 0)
        except (FileNotFoundError, ProcessLookupError, ValueError):
            # A terminated worker can leave its pid marker behind.  Remove the
            # marker before removing the directory so the task can be resumed.
            for child in lock.iterdir():
                child.unlink()
            os.rmdir(lock)
            lock.mkdir()
        else:
            print("SKIP active lock {}".format(lock), flush=True)
            return True
    (lock / "pid").write_text(str(os.getpid()))

    dataset_roots = {
        "cifar224": DATA_ROOT,
        "cub": DATA_ROOT / "cub",
        "imagenetr": DATA_ROOT / "imagenet-r",
        "imagenet_r": DATA_ROOT / "imagenet-r",
        "omnibenchmark": DATA_ROOT / "omnibenchmark",
        "vtab": DATA_ROOT / "vtab",
    }
    args["data_path"] = str(dataset_roots[task["data_key"]])
    args["device"] = [str(gpu_id)]
    args["seed"] = [task["seed"]]
    print("START sheet={} method={} dataset={} seed={} gpu={}".format(task["sheet"], task["method"], task["dataset"], task["seed"], gpu_id), flush=True)
    try:
        try:
            train(args)
        except Exception:
            print(
                "FAILED sheet={} method={} dataset={} seed={}".format(
                    task["sheet"], task["method"], task["dataset"], task["seed"]
                ),
                flush=True,
            )
            traceback.print_exc()
            return False
        success = is_complete(output, expected_tasks)
        print("{} {}".format("COMPLETE" if success else "INCOMPLETE", output), flush=True)
        return success
    finally:
        for child in lock.iterdir():
            child.unlink()
        lock.rmdir()


def main():
    if len(sys.argv) not in (4, 5):
        raise SystemExit(
            "Usage: run_missing_baselines.py <gpu-id> <worker-slot> <worker-count> [--dynamic]"
        )
    gpu_id, slot, worker_count = map(int, sys.argv[1:4])
    dynamic = len(sys.argv) == 5 and sys.argv[4] == "--dynamic"
    if len(sys.argv) == 5 and not dynamic:
        raise SystemExit("The only optional mode is --dynamic.")
    if gpu_id in (0, 3) or slot < 0 or worker_count < 1 or slot >= worker_count:
        raise SystemExit("GPU 0 and 3 are forbidden; worker slot must be in range.")
    tasks = build_tasks()
    assigned = tasks if dynamic else [task for index, task in enumerate(tasks) if index % worker_count == slot]
    print(
        "worker={}/{} gpu={} mode={} assigned_tasks={} total_missing={}".format(
            slot, worker_count, gpu_id, "dynamic" if dynamic else "partitioned", len(assigned), len(tasks)
        ),
        flush=True,
    )
    failures = 0
    for task in assigned:
        if not run_task(task, gpu_id):
            failures += 1
    raise SystemExit(1 if failures else 0)


if __name__ == "__main__":
    main()
