#!/usr/bin/env python
"""Map every CIL result-cell pair to its valid second-round A800 log.

The workbook headers define the number of tasks.  For a B0 experiment the
filename suffix is therefore derived from `total_classes / task_count`, not
from the literal task count: CIFAR=0_10, CUB=0_20, ImageNet-R=0_40,
OmniBenchmark=0_30, VTAB=0_10.
"""

from __future__ import print_function

import ast
import csv
import glob
import os
import re

from openpyxl import load_workbook


WORKBOOK = "CIL_Results.xlsx"
LOG_ROOT = os.path.join("logs", "第二轮A800")
CSV_REPORT = "result_log_mapping.csv"

DATASETS = [
    # (dataset display name, workbook columns, filename-prefix aliases, classes, T)
    ("CIFAR100", ("C", "D"), ("cifar224",), 100, 10),
    ("CUB200", ("E", "F"), ("cub",), 200, 10),
    ("ImageNet-R", ("G", "H"), ("imagenetr", "imagenet_r"), 200, 5),
    ("OmniBenchmark", ("I", "J"), ("omnibenchmark",), 300, 10),
    ("VTAB", ("K", "L"), ("vtab",), 50, 5),
]

METHOD_FOLDERS = {
    "Full Fine-Tuning": "Finetune",
    "SimpleCIL": "SimpleCIL",
    "APER": "adam_adapter",
    "L2P": "L2P",
    "DualPrompt": "DualPrompt",
    "CODA-Prompt": "CODA-Prompt",
    "LAE": "LAE",
    "InfLoRA": "InfLoRA",
    "EASE": "EASE",
    "BiLoRA": "BiLoRA",
    "SD-LoRA": "SD-LoRA-IN21K",
    "CL-LoRA": "CL-LoRA",
    "iCaRL": "iCaRL",
    "FOSTER": "FOSTER",
}

CURVE_RE = re.compile(r"CNN top1 curve: (\[[^\r\n]+\])")
AVERAGE_RE = re.compile(r"Average Accuracy \(CNN\): ([0-9.]+)")


def parse_log(path, expected_tasks):
    """Return status and final CNN curve values for one candidate log."""
    with open(path, "r", errors="ignore") as handle:
        text = handle.read()

    curves = CURVE_RE.findall(text)
    averages = AVERAGE_RE.findall(text)
    if not curves:
        return {"status": "Missing CNN curve", "al": "", "average": "", "curve_len": 0}

    try:
        curve = ast.literal_eval(curves[-1])
    except (SyntaxError, ValueError):
        return {"status": "Unparseable CNN curve", "al": "", "average": "", "curve_len": 0}

    curve_len = len(curve)
    if curve_len != expected_tasks:
        status = "Incomplete: {0}/{1} tasks".format(curve_len, expected_tasks)
    else:
        status = "Complete"
    return {
        "status": status,
        "al": curve[-1] if curve else "",
        "average": averages[-1] if averages else "",
        "curve_len": curve_len,
    }


def compact_value(value):
    return "" if value is None else str(value)


def find_candidates(folder, prefixes, init, increment, seed):
    suffix = "_{0}_{1}_{2}.log".format(init, increment, seed).lower()
    candidates = []
    for path in glob.glob(os.path.join(LOG_ROOT, folder, "*.log")):
        basename = os.path.basename(path).lower()
        if basename.endswith(suffix) and any(basename.startswith(prefix) for prefix in prefixes):
            candidates.append(path)
    return sorted(candidates)


def compare_excel_to_log(excel_al, excel_average, details):
    if len(details) != 1 or details[0]["status"] != "Complete":
        return "Log is incomplete; cannot verify"
    detail = details[0]
    if excel_al == "" and excel_average == "":
        return "Workbook values missing; can be filled from log"
    try:
        same_al = abs(float(excel_al) - float(detail["al"])) < 0.005
        same_average = abs(float(excel_average) - float(detail["average"])) < 0.005
    except (TypeError, ValueError):
        return "Workbook and log values are not comparable"
    return "Match" if same_al and same_average else "Mismatch"


def main():
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    rows = []

    for sheet_name in ("seed0", "seed42", "seed1993", "seed1994"):
        seed = int(sheet_name.replace("seed", ""))
        sheet = workbook[sheet_name]
        for row_number in range(3, sheet.max_row + 1):
            method = sheet.cell(row_number, 1).value
            if method not in METHOD_FOLDERS:
                continue
            folder = METHOD_FOLDERS[method]
            for dataset, columns, prefixes, class_count, task_count in DATASETS:
                init = 0
                increment = class_count // task_count
                candidates = find_candidates(folder, prefixes, init, increment, seed)
                log_details = [parse_log(path, task_count) for path in candidates]

                if not candidates:
                    status = "No matching log in method directory"
                elif len(candidates) == 1:
                    status = log_details[0]["status"]
                else:
                    status = "Multiple candidates: " + "; ".join(
                        detail["status"] for detail in log_details
                    )

                rows.append({
                    "sheet": sheet_name,
                    "seed": seed,
                    "method": method,
                    "dataset": dataset,
                    "cells": "{0}{1}:{2}{1}".format(columns[0], row_number, columns[1]),
                    "excel_al": compact_value(sheet["{0}{1}".format(columns[0], row_number)].value),
                    "excel_average": compact_value(sheet["{0}{1}".format(columns[1], row_number)].value),
                    "expected": "{0}|{1}_{2}_{3}.log".format("/".join(prefixes), init, increment, seed),
                    "schedule": "T={0}; B0/Inc{1} (filename _0_{1}_{2})".format(task_count, increment, seed),
                    "logs": [os.path.relpath(path) for path in candidates],
                    "details": log_details,
                    "status": status,
                    "comparison": compare_excel_to_log(
                        compact_value(sheet["{0}{1}".format(columns[0], row_number)].value),
                        compact_value(sheet["{0}{1}".format(columns[1], row_number)].value),
                        log_details,
                    ),
                })

    # This sheet is an ImageNet-R task-count ablation and does not identify a
    # seed. Only its non-empty result pairs are included; all second-round A800
    # candidates are shown instead of inferring a seed.
    ablation_sheet = workbook["ImageNet-R"]
    ablation_columns = [(("C", "D"), 5), (("E", "F"), 10), (("G", "H"), 20), (("I", "J"), 40)]
    for row_number in range(3, ablation_sheet.max_row + 1):
        method = ablation_sheet.cell(row_number, 1).value
        if method not in METHOD_FOLDERS:
            continue
        folder = METHOD_FOLDERS[method]
        for columns, task_count in ablation_columns:
            excel_al = compact_value(ablation_sheet["{0}{1}".format(columns[0], row_number)].value)
            excel_average = compact_value(ablation_sheet["{0}{1}".format(columns[1], row_number)].value)
            if excel_al == "" and excel_average == "":
                continue
            init = 0
            increment = 200 // task_count
            candidates = []
            for seed in (0, 42, 1993, 1994):
                candidates.extend(find_candidates(folder, ("imagenetr", "imagenet_r"), init, increment, seed))
            candidates = sorted(candidates)
            log_details = [parse_log(path, task_count) for path in candidates]
            if not candidates:
                status = "No matching log in method directory"
            elif len(candidates) == 1:
                status = log_details[0]["status"]
            else:
                status = "Multiple candidates: " + "; ".join(detail["status"] for detail in log_details)
            rows.append({
                "sheet": "ImageNet-R",
                "seed": "not specified",
                "method": method,
                "dataset": "ImageNet-R T={0}".format(task_count),
                "cells": "{0}{1}:{2}{1}".format(columns[0], row_number, columns[1]),
                "excel_al": excel_al,
                "excel_average": excel_average,
                "expected": "imagenetr/imagenet_r_0_{0}_<seed>.log".format(increment),
                "schedule": "T={0}; B0/Inc{1} (seed not specified by this sheet)".format(task_count, increment),
                "logs": [os.path.relpath(path) for path in candidates],
                "details": log_details,
                "status": status,
                "comparison": "Seed is not specified by this sheet; no unique mapping is possible",
            })

    with open(CSV_REPORT, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow([
            "workbook_sheet", "class_order_seed", "method", "dataset", "workbook_cells",
            "workbook_A_L", "workbook_A_avg", "expected_log_filename", "protocol",
            "candidate_log_paths", "log_A_L", "log_A_avg", "log_status",
            "workbook_log_comparison",
        ])
        for row in rows:
            writer.writerow([
                row["sheet"], row["seed"], row["method"], row["dataset"], row["cells"],
                row["excel_al"], row["excel_average"], row["expected"], row["schedule"],
                " ; ".join(row["logs"]),
                " ; ".join(str(detail["al"]) for detail in row["details"]),
                " ; ".join(str(detail["average"]) for detail in row["details"]),
                row["status"], row["comparison"],
            ])

    summary = {
        "complete": 0,
        "incomplete": 0,
        "missing_log": 0,
        "missing_curve": 0,
        "multiple_candidates": 0,
    }
    for row in rows:
        if row["status"] == "Complete":
            summary["complete"] += 1
        elif row["status"].startswith("Incomplete"):
            summary["incomplete"] += 1
        elif row["status"] == "No matching log in method directory":
            summary["missing_log"] += 1
        elif row["status"] == "Missing CNN curve":
            summary["missing_curve"] += 1
        elif row["status"].startswith("Multiple candidates"):
            summary["multiple_candidates"] += 1

    print("wrote", CSV_REPORT)
    print("summary", summary)


if __name__ == "__main__":
    main()
